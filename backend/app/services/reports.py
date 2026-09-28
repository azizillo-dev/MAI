"""Guruh hisoboti: Excel (.xlsx) va PDF.

Ma'lumot bir marta yig'iladi (`group_report_data`), keyin ikki formatga chiqariladi.
Fayl yaratish CPU talab qiladi — chaqiruvchi uni alohida oqimda bajaradi.
"""

import io
from datetime import datetime, timedelta
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import as_utc, utcnow
from app.models import Assignment, AssignmentStatus, Group, GroupMember, MemberStatus, Submission, SubmissionStatus, User
from app.services.assignments import TASHKENT

FONTS = Path(__file__).resolve().parent.parent / "assets" / "fonts"
SUBJECTS = {"math": "Matematika", "english": "Ingliz tili"}
MONTHS = ["yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul", "avgust", "sentabr", "oktabr", "noyabr",
          "dekabr"]


def _grade(percent: float, scale: str) -> float:
    from app.services.assignments import to_scale

    return to_scale(percent, scale)


async def group_report_data(db: AsyncSession, group: Group, period: str) -> dict:
    now = utcnow()
    since = now - timedelta(days=30) if period == "month" else None
    members = list(await db.scalars(
        select(User).join(GroupMember, GroupMember.student_id == User.id)
        .where(GroupMember.group_id == group.id, GroupMember.status == MemberStatus.ACTIVE)
        .order_by(User.last_name, User.first_name)
    ))
    q = select(Assignment.id, Assignment.title, Assignment.due_at).where(
        Assignment.group_id == group.id, Assignment.status == AssignmentStatus.PUBLISHED)
    if since is not None:
        q = q.where(Assignment.due_at >= since)
    assignments = (await db.execute(q.order_by(Assignment.due_at))).all()
    aids = [a.id for a in assignments]
    subs = (await db.execute(
        select(Submission.assignment_id, Submission.student_id, Submission.status, Submission.final_score,
               Submission.is_late)
        .where(Submission.assignment_id.in_(aids))
    )).all() if aids else []

    scale = float(group.grading_scale)
    due_ids = {a.id for a in assignments if as_utc(a.due_at) < now}
    by = {(s.student_id, s.assignment_id): s for s in subs}

    def pct(s) -> float | None:
        return s.final_score / scale * 100 if s and s.status == SubmissionStatus.GRADED and s.final_score is not None else None

    students = []
    for m in members:
        mine = [by.get((m.id, a.id)) for a in assignments]
        ps = [p for s in mine if (p := pct(s)) is not None]
        avg = sum(ps) / len(ps) if ps else None
        students.append({
            "name": m.full_name,
            "submitted": sum(1 for s in mine if s is not None),
            "missed": sum(1 for a in assignments if a.id in due_ids and (m.id, a.id) not in by),
            "late": sum(1 for s in mine if s is not None and s.is_late),
            "avg_percent": round(avg, 1) if avg is not None else None,
            "grade": _grade(avg, group.grading_scale) if avg is not None else None,
        })
    rows_a = []
    for a in assignments:
        a_subs = [s for s in subs if s.assignment_id == a.id]
        ps = [p for s in a_subs if (p := pct(s)) is not None]
        rows_a.append({"title": a.title, "due": as_utc(a.due_at).astimezone(TASHKENT), "submitted": len(a_subs),
                       "avg_percent": round(sum(ps) / len(ps), 1) if ps else None})
    all_p = [s["avg_percent"] for s in students if s["avg_percent"] is not None]
    local_now = now.astimezone(TASHKENT)
    return {
        "group": group.name,
        "subject": SUBJECTS.get(group.subject, group.subject),
        "scale": group.grading_scale,
        "period": "Oxirgi 30 kun" if since else "Butun davr",
        "generated": f"{local_now.day}-{MONTHS[local_now.month - 1]} {local_now.year}, {local_now:%H:%M}",
        "students": students,
        "assignments": rows_a,
        "summary": {
            "students": len(members),
            "assignments": len(assignments),
            "avg_percent": round(sum(all_p) / len(all_p), 1) if all_p else None,
            "low": sum(1 for p in all_p if p < 60),
        },
    }


def _fmt(v, suffix: str = "") -> str:
    if v is None:
        return "—"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return f"{v}{suffix}"


def build_xlsx(d: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "O'quvchilar"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4F46E5")
    ws.append([f"{d['group']} — {d['subject']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{d['period']} · {d['summary']['students']} o'quvchi · {d['summary']['assignments']} vazifa · "
               f"o'rtacha {_fmt(d['summary']['avg_percent'], '%')} · tayyorlandi: {d['generated']}"])
    ws.append([])
    cols = ["O'quvchi", "Topshirgan", "Topshirmagan", "Kechikkan", "O'rtacha, %", f"Baho ({d['scale']} ballik)"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=4, column=c)
        cell.font, cell.fill, cell.alignment = head, fill, Alignment(horizontal="center")
    for s in d["students"]:
        ws.append([s["name"], s["submitted"], s["missed"], s["late"], s["avg_percent"], s["grade"]])
    for i, w in enumerate([30, 12, 14, 12, 13, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    wa = wb.create_sheet("Vazifalar")
    wa.append(["Vazifa", "Muddat", "Topshirgan", "O'rtacha, %"])
    for c in range(1, 5):
        cell = wa.cell(row=1, column=c)
        cell.font, cell.fill = head, fill
    for a in d["assignments"]:
        wa.append([a["title"], a["due"].strftime("%d.%m.%Y %H:%M"), a["submitted"], a["avg_percent"]])
    for i, w in enumerate([40, 18, 12, 13], start=1):
        wa.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(d: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_font("DejaVu", "", str(FONTS / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(FONTS / "DejaVuSans-Bold.ttf"))
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()

    pdf.set_fill_color(79, 70, 229)
    pdf.rect(0, 0, 210, 34, "F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("DejaVu", "B", 18)
    pdf.set_xy(12, 9)
    pdf.cell(0, 9, f"{d['group']} — guruh hisoboti")
    pdf.set_font("DejaVu", "", 10)
    pdf.set_xy(12, 20)
    pdf.cell(0, 6, f"Mentor AI · {d['subject']} · {d['period']} · {d['generated']}")
    pdf.set_text_color(15, 23, 42)
    pdf.set_y(42)

    # Qisqa ko'rsatkichlar kartalari
    sm = d["summary"]
    cards = [("O'quvchilar", _fmt(sm["students"])), ("Vazifalar", _fmt(sm["assignments"])),
             ("O'rtacha", _fmt(sm["avg_percent"], "%")), ("60% dan past", _fmt(sm["low"]))]
    x = 12
    for label, value in cards:
        pdf.set_fill_color(241, 245, 249)
        pdf.rect(x, 42, 44, 20, "F")
        pdf.set_xy(x + 4, 45)
        pdf.set_font("DejaVu", "B", 14)
        pdf.cell(36, 7, value)
        pdf.set_xy(x + 4, 53)
        pdf.set_font("DejaVu", "", 8)
        pdf.cell(36, 5, label)
        x += 47
    pdf.set_y(70)

    def table(title: str, headers: list[str], widths: list[int], rows: list[list[str]]) -> None:
        pdf.set_font("DejaVu", "B", 12)
        pdf.cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "B", 9)
        pdf.set_fill_color(79, 70, 229)
        pdf.set_text_color(255, 255, 255)
        for h, w in zip(headers, widths, strict=True):
            pdf.cell(w, 8, h, fill=True, align="C")
        pdf.ln()
        pdf.set_text_color(15, 23, 42)
        pdf.set_font("DejaVu", "", 9)
        for i, row in enumerate(rows):
            pdf.set_fill_color(248, 250, 252) if i % 2 else pdf.set_fill_color(255, 255, 255)
            for j, (v, w) in enumerate(zip(row, widths, strict=True)):
                pdf.cell(w, 7, v, fill=True, align="L" if j == 0 else "C")
            pdf.ln()
        pdf.ln(4)

    table("O'quvchilar", ["O'quvchi", "Topshirgan", "Qoldirgan", "Kechikkan", "O'rtacha", "Baho"],
          [66, 24, 24, 24, 24, 24],
          [[s["name"][:34], _fmt(s["submitted"]), _fmt(s["missed"]), _fmt(s["late"]), _fmt(s["avg_percent"], "%"),
            _fmt(s["grade"])] for s in d["students"]] or [["O'quvchi yo'q", "", "", "", "", ""]])
    table("Vazifalar", ["Vazifa", "Muddat", "Topshirgan", "O'rtacha"], [96, 34, 28, 28],
          [[a["title"][:50], a["due"].strftime("%d.%m.%Y"), _fmt(a["submitted"]), _fmt(a["avg_percent"], "%")]
           for a in d["assignments"]] or [["Bu davrda vazifa yo'q", "", "", ""]])
    return bytes(pdf.output())


def filename(d: dict, ext: str) -> str:
    safe = "".join(ch if ch.isalnum() else "_" for ch in d["group"]).strip("_") or "guruh"
    return f"{safe}_hisobot_{datetime.now(TASHKENT):%Y-%m-%d}.{ext}"
